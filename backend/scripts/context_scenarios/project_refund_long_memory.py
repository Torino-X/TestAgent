"""Real API + real model project-memory scenario: CTX-PROJECT-01.

This runner is intentionally not a pytest wrapper and never writes to the
database directly.  It creates a disposable project, generates three real
DOCX documents, uploads them through the product API, then drives meaningful
normal-chat turns through the configured model.  The final recall checks use
deterministic fact atoms, not an LLM judge.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import os
import re
import sys
import tempfile
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

SCRIPTS_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = SCRIPTS_ROOT.parent
for path in (SCRIPTS_ROOT, BACKEND_ROOT):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from phase2_lct_runtime import LctRun, cli_failure, repository_root
from phase2_live_api_e2e import ApiClient, ApiFailure, _data, _request_id


SCENARIO = "CTX-PROJECT-01"
PROJECT_TITLE_PREFIX = "CTX-PROJECT-01 退款功能上线测试 "
CONVERSATION_TITLE_PREFIX = "CTX-PROJECT-01 "
FILE_NAME_PREFIX = "CTX_PROJECT_01_"
STATE_NAME = "CTX-PROJECT-01-latest-passed-project.json"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ScenarioStop(RuntimeError):
    """A terminal boundary where continuing would make API state ambiguous."""


class MultipartApiClient(ApiClient):
    """Small extension of the existing credential-safe JSON client for uploads."""

    def upload(self, path: str, *, filename: str, content: bytes, fields: dict[str, str]) -> Any:
        boundary = f"----ctxproject{uuid.uuid4().hex}"
        parts: list[bytes] = []
        for name, value in fields.items():
            parts.extend([
                f"--{boundary}\r\n".encode(),
                f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                value.encode("utf-8"),
                b"\r\n",
            ])
        media_type = mimetypes.guess_type(filename)[0] or DOCX_MIME
        parts.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode(),
            f"Content-Type: {media_type}\r\n\r\n".encode(),
            content,
            b"\r\n",
            f"--{boundary}--\r\n".encode(),
        ])
        headers = {
            "Accept": "application/json",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        }
        if self.token:
            headers["Authorization"] = f"Bearer {self.token}"
        request = urllib.request.Request(
            self.base_url + path,
            data=b"".join(parts),
            method="POST",
            headers=headers,
        )
        try:
            with self.opener.open(request, timeout=self.timeout_seconds) as response:
                raw = response.read().decode("utf-8")
                decoded = json.loads(raw) if raw else None
                self.calls.append({"method": "POST", "path": path, "status": response.status, "request_id": _request_id(decoded)})
                return decoded
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", errors="replace")
            try:
                decoded = json.loads(raw)
            except json.JSONDecodeError:
                decoded = {"non_json_body": raw[:500]}
            self.calls.append({"method": "POST", "path": path, "status": exc.code, "request_id": _request_id(decoded)})
            raise ApiFailure("POST", path, exc.code, decoded) from exc


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def _short_hash(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:16]


def _normalise(value: str) -> str:
    return re.sub(r"[\s\-_/：:，,。；;、（）()【】\[\]‘’'\".!！?？]+", "", value).lower()


def _atom_result(reply: str, required_groups: Iterable[Iterable[str]], forbidden: Iterable[str] = ()) -> dict[str, Any]:
    normalized = _normalise(reply)
    missing = [list(group) for group in required_groups if not any(_normalise(atom) in normalized for atom in group)]
    forbidden_hits = [atom for atom in forbidden if _normalise(atom) in normalized]
    return {
        "passed": not missing and not forbidden_hits,
        "missing_any_of_groups": missing,
        "forbidden_hits": forbidden_hits,
        "reply_length": len(reply),
        "reply_sha256_16": _short_hash(reply),
        # All prompts/replies in this scenario are test-owned refund material.
        # A bounded excerpt materially improves failure diagnosis after a failed
        # run's project is cleaned up; credentials/tokens are never captured.
        "reply_excerpt": reply[:1200],
    }


def _reply_text(response: Any) -> str:
    if not isinstance(response, dict):
        return ""
    reply = response.get("agent_reply")
    return str(reply.get("content") or "") if isinstance(reply, dict) else ""


def _safe_context_diagnostics(response: Any) -> dict[str, Any]:
    reply = response.get("agent_reply") if isinstance(response, dict) else None
    payload = reply.get("payload") if isinstance(reply, dict) else None
    diagnostics = payload.get("chat_context_engine") if isinstance(payload, dict) else None
    if not isinstance(diagnostics, dict):
        return {}
    allowed = ("path", "outcome", "snapshot_public_id", "exception_type", "error_code", "stage", "retryable")
    return {key: diagnostics[key] for key in allowed if key in diagnostics}


def _safe_failure(exc: Exception, *, method: str, path: str, timeout_seconds: int) -> dict[str, Any]:
    if isinstance(exc, ApiFailure):
        body = exc.body if isinstance(exc.body, dict) else {}
        detail = next((str(body[key]) for key in ("detail", "message", "error", "code") if body.get(key) is not None), "")
        detail = re.sub(r"(?i)(password|token|authorization)\s*[:=]\s*[^\s,;]+", r"\1=[REDACTED]", detail)
        return {
            "method": exc.method,
            "path": exc.path,
            "status": exc.status,
            "response_keys": sorted(str(key) for key in body),
            "response_detail": detail[:500] or None,
        }
    return {
        "method": method,
        "path": path,
        "status": None,
        "failure_type": type(exc).__name__,
        "timeout_seconds": timeout_seconds,
    }


def _require_id(data: Any, *, keys: Iterable[str], operation: str) -> str:
    if isinstance(data, dict):
        for key in keys:
            value = data.get(key)
            if value:
                return str(value)
    raise RuntimeError(f"{operation} response contained none of: {', '.join(keys)}")


def _is_project_public_id(value: str) -> bool:
    """Match the production Project API's stable public-ID prefix."""
    return str(value).startswith("prj_")


def _make_fixture_docs(directory: Path, marker: str) -> list[dict[str, Any]]:
    """Generate actual OOXML DOCX files; no fake text payload is uploaded."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - execution environment contract
        raise RuntimeError("python-docx is required to create CTX-PROJECT-01 fixture documents") from exc

    documents = [
        {
            "label": "PRD",
            "filename": f"{FILE_NAME_PREFIX}{marker}_退款功能_PRD_v1.2.docx",
            "role": "requirement",
            "title": "退款功能 PRD v1.2",
            "sections": [
                ("1. 适用订单", ["本期自助退款仅适用于订单状态为“已签收”或“已完成”的订单。"]),
                ("2. 时限规则", ["用户必须在签收时间起 72 小时内提交退款申请。", "超过签收后 72 小时的申请不予受理，前端应提示用户联系人工客服。"]),
                ("3. 物流争议", ["物流争议不可自动退款，必须转人工处理并记录争议原因。"]),
            ],
        },
        {
            "label": "API",
            "filename": f"{FILE_NAME_PREFIX}{marker}_退款开放平台_API_v1.2.docx",
            "role": "api_spec",
            "title": "退款开放平台 API v1.2",
            "sections": [
                ("1. 创建退款", ["接口：POST /refunds。", "请求字段 requestId 为调用方生成的幂等键。"]),
                ("2. 幂等规则", ["同一 requestId 在 24 小时内只能成功创建一次退款申请。", "在有效期内重复提交相同 requestId 时，接口返回错误码 RF409。"]),
                ("3. 返回处理", ["客户端收到 RF409 后不得再次创建退款，应查询原退款申请状态。"]),
            ],
        },
        {
            "label": "UAT会议纪要",
            "filename": f"{FILE_NAME_PREFIX}{marker}_退款功能_UAT评审纪要.docx",
            "role": "historical_test",
            "title": "退款功能 UAT 评审纪要",
            "sections": [
                ("1. 评审建议", ["评审阶段建议关注支付失败退款、重复退款和签收后时限边界。", "物流争议和人工审核时效因用户影响较大，建议在范围确认会上进一步讨论。"]),
                ("2. 结论边界", ["本纪要记录的是评审建议和待确认事项，不构成最终 UAT 范围或验收负责人任命。"]),
                ("3. 发布背景", ["目标上线窗口为 2026 年 10 月；UAT 范围以负责人在项目会话中的最终确认作为准则。"]),
            ],
        },
    ]
    for spec in documents:
        document = Document()
        document.add_heading(spec["title"], level=0)
        document.add_paragraph("CTX-PROJECT-01 自动化场景资料。内容仅用于退款功能上线测试。")
        for heading, paragraphs in spec["sections"]:
            document.add_heading(heading, level=1)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph, style="List Bullet")
        path = directory / spec["filename"]
        document.save(path)
        spec["path"] = path
        spec["sha256_16"] = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    return documents


def _state_path() -> Path:
    return repository_root() / "test-results" / "phase2-resource-pack" / STATE_NAME


def _read_state() -> dict[str, Any] | None:
    try:
        value = json.loads(_state_path().read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    return value if isinstance(value, dict) and value.get("scenario") == SCENARIO else None


def _write_state(*, project_id: str, conversation_id: str, file_ids: list[str]) -> None:
    path = _state_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({
        "scenario": SCENARIO,
        "project_id": project_id,
        "conversation_id": conversation_id,
        "file_ids": file_ids,
        "retained_at_utc": datetime.now(timezone.utc).isoformat(),
    }, ensure_ascii=False, indent=2), encoding="utf-8")


def _remove_state() -> None:
    try:
        _state_path().unlink()
    except FileNotFoundError:
        pass


def _owned_previous_state(client: ApiClient, state: dict[str, Any]) -> bool:
    project_id = state.get("project_id")
    conversation_id = state.get("conversation_id")
    file_ids = state.get("file_ids")
    if not isinstance(project_id, str) or not isinstance(conversation_id, str) or not isinstance(file_ids, list) or len(file_ids) != 3:
        return False
    try:
        project = _data(client.request("GET", f"/projects/{urllib.parse.quote(project_id)}"))
        conversation = _data(client.request("GET", f"/conversations/{urllib.parse.quote(conversation_id)}"))
        sources = _data(client.request("GET", f"/projects/{urllib.parse.quote(project_id)}/sources"))
    except (ApiFailure, OSError, TimeoutError):
        return False
    items = sources.get("items", []) if isinstance(sources, dict) else []
    source_file_ids = {str(item.get("fileId")) for item in items if isinstance(item, dict)}
    return (
        isinstance(project, dict)
        and str(project.get("name") or "").startswith(PROJECT_TITLE_PREFIX)
        and isinstance(conversation, dict)
        and str(conversation.get("title") or "").startswith(CONVERSATION_TITLE_PREFIX)
        and set(map(str, file_ids)).issubset(source_file_ids)
    )


def _delete_records(client: ApiClient, *, project_id: str | None, conversation_ids: Iterable[str], file_ids: Iterable[str]) -> list[dict[str, Any]]:
    """Delete only exact IDs created by this runner, retaining diagnostic outcomes."""
    outcomes: list[dict[str, Any]] = []
    for conversation_id in dict.fromkeys(conversation_ids):
        if not conversation_id:
            continue
        try:
            _data(client.request("DELETE", f"/conversations/{urllib.parse.quote(conversation_id)}"))
            outcomes.append({"kind": "conversation", "id": conversation_id, "deleted": True})
        except Exception as exc:  # cleanup cannot hide primary evidence
            outcomes.append({"kind": "conversation", "id": conversation_id, "deleted": False, "error": type(exc).__name__})
    if project_id:
        try:
            _data(client.request("DELETE", f"/projects/{urllib.parse.quote(project_id)}"))
            outcomes.append({"kind": "project", "id": project_id, "deleted": True})
        except Exception as exc:
            outcomes.append({"kind": "project", "id": project_id, "deleted": False, "error": type(exc).__name__})
    for file_id in dict.fromkeys(file_ids):
        if not file_id:
            continue
        try:
            # The files were uploaded by this exact runner and are being
            # removed only after their project binding has been deleted.
            # The product's permanent endpoint intentionally accepts only a
            # soft-deleted item.  Keep the two calls adjacent and scoped to
            # the exact IDs created by this runner.
            _data(client.request("DELETE", f"/library/items/{urllib.parse.quote(file_id)}"))
            _data(client.request("DELETE", f"/library/items/{urllib.parse.quote(file_id)}/permanent"))
            outcomes.append({"kind": "library_file", "id": file_id, "deleted": True})
        except Exception as exc:
            outcomes.append({"kind": "library_file", "id": file_id, "deleted": False, "error": type(exc).__name__})
    return outcomes


def _index_statuses(client: ApiClient, file_ids: Iterable[str]) -> list[dict[str, Any]]:
    """Read only owner-scoped index state for the three files just uploaded."""
    file_id_set = {str(file_id) for file_id in file_ids}
    rows = _data(client.request("GET", "/context/index/documents?limit=100"))
    if not isinstance(rows, list):
        return []
    return [
        {
            "document_public_id": row.get("document_public_id"),
            "source_public_id": row.get("source_public_id"),
            "status": row.get("status"),
            "lexical_index_status": row.get("lexical_index_status"),
            "vector_index_status": row.get("vector_index_status"),
        }
        for row in rows
        if isinstance(row, dict) and str(row.get("source_public_id")) in file_id_set
    ]


def _wait_for_indexing(client: ApiClient, *, file_ids: list[str], wait_seconds: int) -> tuple[bool, list[dict[str, Any]]]:
    """Avoid classifying an asynchronous indexing delay as a memory failure."""
    if len(file_ids) != 3:
        return False, []
    deadline = time.monotonic() + wait_seconds
    latest: list[dict[str, Any]] = []
    while True:
        latest = _index_statuses(client, file_ids)
        by_source = {str(item.get("source_public_id")): item for item in latest}
        all_seen = len(by_source) == len(file_ids)
        all_indexed = all_seen and all(str(by_source[file_id].get("status")) == "indexed" for file_id in file_ids)
        if all_indexed:
            return True, latest
        terminal_failure = any(str(item.get("status")) == "failed" for item in latest)
        if terminal_failure or time.monotonic() >= deadline:
            return False, latest
        time.sleep(3)


class ProjectScenario:
    def __init__(self, *, run: LctRun, client: MultipartApiClient, project_id: str, conversation_id: str, request_timeout_seconds: int) -> None:
        self.run = run
        self.client = client
        self.project_id = project_id
        self.conversation_id = conversation_id
        self.request_timeout_seconds = request_timeout_seconds
        self.turns: list[dict[str, Any]] = []

    def send(self, *, label: str, prompt: str, oracle: dict[str, Any] | None = None) -> dict[str, Any]:
        path = f"/conversations/{urllib.parse.quote(self.conversation_id)}/messages"
        try:
            response = _data(self.client.request("POST", path, {
                "content": prompt,
                "attached_file_ids": [],
                "knowledge_mode_snapshot": "AUTO",
            }))
        except (ApiFailure, OSError, TimeoutError) as exc:
            diagnostic = _safe_failure(exc, method="POST", path=path, timeout_seconds=self.request_timeout_seconds)
            record = {"label": label, "api_failure": diagnostic, "response_received": False}
            self.turns.append(record)
            step = self.run.step(f"{label}: normal-chat API response", component="Live normal chat API", expected="successful completed response")
            self.run.observe(step, False, actual=json.dumps(diagnostic, ensure_ascii=False), evidence={"turn": record})
            # A transport timeout leaves persistence state unknown, so later
            # turns could accidentally create an invalid transcript.
            if not isinstance(exc, ApiFailure):
                raise ScenarioStop(f"ambiguous normal-chat transport failure at {label}")
            return record

        reply = _reply_text(response)
        record = {
            "label": label,
            "response_received": True,
            "route": response.get("route") if isinstance(response, dict) else None,
            "intent": response.get("intent") if isinstance(response, dict) else None,
            "requires_sse": response.get("requires_sse") if isinstance(response, dict) else None,
            "chat_context_diagnostics": _safe_context_diagnostics(response),
            "reply_length": len(reply),
            "reply_sha256_16": _short_hash(reply),
        }
        self.turns.append(record)
        route_step = self.run.step(f"{label}: normal-chat route", component="Intent routing + Context Engine", expected="route=chat_reply")
        self.run.observe(route_step, record["route"] == "chat_reply", actual=json.dumps({"route": record["route"], "intent": record["intent"]}, ensure_ascii=False), evidence={"turn": record})
        if oracle:
            atom = _atom_result(reply, oracle["required_groups"], oracle.get("forbidden", ()))
            record["oracle"] = atom
            oracle_step = self.run.step(oracle["name"], component=oracle["component"], expected=oracle["expected"])
            self.run.observe(oracle_step, bool(atom["passed"]), actual=json.dumps(atom, ensure_ascii=False), evidence={"turn": record})
        return record


BASELINE_PROMPTS = [
    (
        "source-prd-deadline",
        "请只基于当前项目资料，说明已签收或已完成订单申请退款的时限，以及超过时限时的处理。不要创建任务或文件，简洁回答。",
        {"name": "PRD 资料事实可回答", "component": "Project source retrieval + real model", "expected": "回答包含 72 小时和超过时限不予受理", "required_groups": (("72小时",), ("不予受理", "不受理", "不能受理"))},
    ),
    (
        "source-api-idempotency",
        "请只基于当前项目资料说明退款接口如何防止重复提交：字段名、有效期和重复时的错误码分别是什么？不要创建任务或文件。",
        {"name": "API 资料事实可回答", "component": "Project source retrieval + real model", "expected": "回答包含 requestId、24 小时和 RF409", "required_groups": (("requestId",), ("24小时",), ("RF409",))},
    ),
    (
        "source-logistics",
        "请只基于当前项目资料说明物流争议退款应如何处理。不要创建任务或文件。",
        {"name": "PRD 物流争议事实可回答", "component": "Project source retrieval + real model", "expected": "回答包含物流争议和人工处理", "required_groups": (("物流争议",), ("人工",))},
    ),
]


WORKING_PROMPTS = [
    "请把签收后 72 小时边界拆成小于、等于、大于三个测试数据点，并说明每个数据点的预期。不要创建任务或文件。",
    "针对同一 requestId 的两次退款请求，解释 UAT 需要分别观察哪些接口结果。不要创建任务或文件。",
    "支付失败退款在本轮验证中最需要关注哪些用户可见结果？只给简短分析。",
    "如果退款接口返回 RF409，测试人员下一步应核对什么，而不是再次提交退款？",
    "将物流争议写成一条发布风险：风险内容、当前处理路径和需要谁关注。不要创建任务或文件。",
    "从评审纪要角度看，哪些事项还属于建议或待确认，而不是已经生效的最终范围？",
    "请给出一条适合日报的 UAT 进度表述，包含已验证、待验证和风险三部分。",
    "对于签收后 72 小时零界点，测试数据的时间精度为什么需要明确？",
    "如果发现重复退款拦截失败，如何描述它对发布的影响？只分析，不创建缺陷或任务。",
    "请比较支付失败退款与物流争议：它们在自动处理路径上有什么不同？",
    "作为验收前检查，列出应再次确认的三个项目资料事实。不要创建任务或文件。",
    "请用两句话概括本轮退款 UAT 的发布风险管理思路。不要创建任务或文件。",
]


DELAYED_PROBES = [
    (
        "delayed-scope-owner",
        "我们早些时候已经明确不纳入本轮 UAT 的两个范围是什么？验收负责人是谁？请直接回答并区分这与评审纪要建议的关系。",
        {"name": "长期会话决策回忆", "component": "Conversation context continuity + real model", "expected": "记住物流争议、人工审核时效、王晨，并区分最终决定与评审建议", "required_groups": (("物流争议",), ("人工审核时效", "人工审核"), ("王晨",), ("建议", "评审纪要"))},
    ),
    (
        "delayed-api-details",
        "重复退款的 API 防重字段是什么、有效多久、重复时怎么处理？请直接回答。",
        {"name": "跨资料长期事实回忆", "component": "Project source retrieval + real model", "expected": "记住 requestId、24 小时和 RF409", "required_groups": (("requestId",), ("24小时",), ("RF409",))},
    ),
    (
        "delayed-deadline-decision",
        "用户签收第 4 天才申请退款，给出可否受理和依据；同时说明这是否属于本轮 P0。",
        {"name": "跨资料与范围决策回忆", "component": "Project source + conversation context + real model", "expected": "超过 72 小时不受理，且 72 小时边界属于 P0", "required_groups": (("72小时",), ("不予受理", "不受理", "不能受理"), ("P0",))},
    ),
    (
        "delayed-p0-matrix",
        "根据 PRD、API 和我们已经确定的范围，列出当前 P0 UAT 用例及其依据；不要把仅记录风险的事项写进 P0。",
        {"name": "综合 P0 用例回忆", "component": "Project source + conversation context + real model", "expected": "P0 包含支付失败、重复退款和 72 小时边界；物流争议明确不纳入", "required_groups": (("支付失败",), ("重复退款",), ("72小时",), ("物流争议",), ("不纳入", "不包含", "排除"))},
    ),
]

# The delayed matrix is a memory/scope question, not a request to generate
# test cases.  The latter is deliberately unsupported by normal chat routing
# and would test product capability gating rather than Context Engine recall.
DELAYED_PROBES = [
    (
        label,
        "根据 PRD、API 和我们已经确认的范围，梳理当前 P0 的范围口径及其依据；请明确哪些内容纳入 P0、哪些内容不纳入 P0，只做范围口径梳理。"
        if label == "delayed-p0-matrix"
        else prompt,
        oracle,
    )
    for label, prompt, oracle in DELAYED_PROBES
]


def _run(args: argparse.Namespace, run: LctRun) -> str:
    username = os.environ.get("PHASE2_E2E_USERNAME", "")
    password = os.environ.get("PHASE2_E2E_PASSWORD", "")
    if not username or not password:
        raise RuntimeError("PHASE2_E2E_USERNAME and PHASE2_E2E_PASSWORD must be supplied by the PowerShell wrapper")

    marker = _stamp()
    client = MultipartApiClient(args.base_url, timeout_seconds=args.request_timeout_seconds)
    project_id: str | None = None
    primary_conversation_id: str | None = None
    negative_conversation_id: str | None = None
    file_ids: list[str] = []
    scenario: ProjectScenario | None = None
    retained = False
    try:
        run.log("CTX-PROJECT-01: authenticate and create the isolated refund-release project")
        auth = run.step("authenticated project E2E session", component="Auth API", expected="login succeeds without storing credentials")
        client.login(username, password)
        me = _data(client.request("GET", "/auth/me"))
        run.check(auth, isinstance(me, dict), actual="authenticated owner-scoped API session")

        project = _data(client.request("POST", "/projects", {
            "name": f"{PROJECT_TITLE_PREFIX}{marker}",
            "description": "自动化真实 API 场景：退款功能上线测试。仅保留最新完整通过样本。",
            "memoryMode": "project_memory",
        }))
        project_id = _require_id(project, keys=("id", "project_id", "public_id"), operation="create project")
        primary = _data(client.request("POST", f"/projects/{urllib.parse.quote(project_id)}/conversations", {"title": f"{CONVERSATION_TITLE_PREFIX}主会话 {marker}"}))
        primary_conversation_id = _require_id(primary, keys=("id", "conversation_id", "public_id"), operation="create project conversation")
        create_step = run.step("project and main conversation created", component="Project API", expected="project-scoped conversation is created")
        run.check(create_step, _is_project_public_id(project_id) and primary_conversation_id.startswith("conv_"), actual=json.dumps({"project_id": project_id, "conversation_id": primary_conversation_id}, ensure_ascii=False))

        with tempfile.TemporaryDirectory(prefix="ctx-project-01-") as fixture_root:
            specs = _make_fixture_docs(Path(fixture_root), marker)
            fixture_manifest = run.write_evidence("fixture-manifest.json", [{key: value for key, value in spec.items() if key != "path"} for spec in specs])
            run.log("CTX-PROJECT-01: upload three real DOCX sources through the project API")
            for spec in specs:
                upload_step = run.step(f"upload {spec['label']} DOCX", component="Project source upload API", expected="actual DOCX is uploaded and bound to this project")
                try:
                    uploaded = _data(client.upload(
                        f"/projects/{urllib.parse.quote(project_id)}/sources/upload",
                        filename=str(spec["filename"]),
                        content=Path(spec["path"]).read_bytes(),
                        fields={"source_role": str(spec["role"])},
                    ))
                    file_id = _require_id(uploaded, keys=("fileId", "file_id"), operation=f"upload {spec['label']} DOCX")
                    file_ids.append(file_id)
                    run.observe(upload_step, bool(uploaded.get("id")) and bool(file_id), actual=json.dumps({"source_id": uploaded.get("id"), "file_id": file_id, "source_role": uploaded.get("sourceRole")}, ensure_ascii=False), evidence={"fixture_manifest": fixture_manifest})
                except Exception as exc:
                    run.observe(upload_step, False, actual=f"{type(exc).__name__}: {exc}", evidence={"fixture_manifest": fixture_manifest})

        sources = _data(client.request("GET", f"/projects/{urllib.parse.quote(project_id)}/sources"))
        source_items = sources.get("items", []) if isinstance(sources, dict) else []
        source_step = run.step("three project sources visible", component="Project source API", expected="three uploaded DOCX sources are current and carry the intended roles")
        role_set = {str(item.get("sourceRole")) for item in source_items if isinstance(item, dict)}
        run.observe(source_step, len(source_items) == 3 and {"requirement", "api_spec", "historical_test"}.issubset(role_set), actual=json.dumps({"count": len(source_items), "roles": sorted(role_set), "file_ids": file_ids}, ensure_ascii=False))
        if len(file_ids) != 3:
            run.notes.append("SOURCE_UPLOAD_INCOMPLETE: source-grounded model turns were skipped because one or more DOCX uploads did not return a project file ID.")
            return "TEST_FAIL"

        run.log("CTX-PROJECT-01: wait for the three project DOCX indexes to become searchable")
        indexes_ready, index_statuses = _wait_for_indexing(client, file_ids=file_ids, wait_seconds=args.index_wait_seconds)
        index_evidence = run.write_evidence("source-index-status.json", {"file_ids": file_ids, "statuses": index_statuses, "ready": indexes_ready})
        index_step = run.step("three project DOCX indexes ready", component="Context Index API", expected="each uploaded source reaches indexed before source-grounded chat begins")
        run.observe(index_step, indexes_ready, actual=json.dumps(index_statuses, ensure_ascii=False), evidence={"index_status_evidence": index_evidence})
        if not indexes_ready:
            run.notes.append("SOURCE_INDEX_NOT_READY: source-grounded model turns were skipped to avoid misclassifying asynchronous indexing as a memory failure.")
            return "TEST_FAIL"

        scenario = ProjectScenario(run=run, client=client, project_id=project_id, conversation_id=primary_conversation_id, request_timeout_seconds=args.request_timeout_seconds)
        run.log("CTX-PROJECT-01: establish source-grounded baseline answers through normal chat")
        for label, prompt, oracle in BASELINE_PROMPTS:
            scenario.send(label=label, prompt=prompt, oracle=oracle)

        run.log("CTX-PROJECT-01: record the user's final UAT scope decision in the real project conversation")
        scenario.send(
            label="user-final-scope-decision",
            prompt="我现在确认：本轮 UAT 的 P0 只覆盖支付失败退款、重复退款拦截、72 小时边界；物流争议和人工审核时效只记录风险，不纳入本轮 UAT；验收负责人是王晨。请确认你会在后续讨论中按这一最终决定执行，不要创建任务或文件。",
        )

        run.log(f"CTX-PROJECT-01: conduct {len(WORKING_PROMPTS)} meaningful project working turns")
        for ordinal, prompt in enumerate(WORKING_PROMPTS, start=1):
            scenario.send(label=f"working-turn-{ordinal}", prompt=prompt)

        run.log("CTX-PROJECT-01: run delayed user-memory probes without re-supplying the facts")
        for label, prompt, oracle in DELAYED_PROBES:
            scenario.send(label=label, prompt=prompt, oracle=oracle)

        run.log("CTX-PROJECT-01: create a fresh project conversation for scope-isolation negative control")
        negative = _data(client.request("POST", f"/projects/{urllib.parse.quote(project_id)}/conversations", {"title": f"{CONVERSATION_TITLE_PREFIX}隔离验证 {marker}"}))
        negative_conversation_id = _require_id(negative, keys=("id", "conversation_id", "public_id"), operation="create negative-control conversation")
        negative_scenario = ProjectScenario(run=run, client=client, project_id=project_id, conversation_id=negative_conversation_id, request_timeout_seconds=args.request_timeout_seconds)
        negative_scenario.send(
            label="fresh-conversation-isolation",
            prompt="这是本项目的一条新会话。项目资料能否确定本轮最终 P0 的排除项和验收负责人？如果不能，请明确说明信息缺失；不要猜测，也不要创建任务或文件。",
            oracle={
                "name": "跨会话用户决策不泄漏",
                "component": "Conversation scope isolation + real model",
                "expected": "新会话明确不知道先前会话的最终决定，且不声称王晨是负责人",
                "required_groups": (("无法", "不能", "未提供", "不清楚", "未知", "无法确定"),),
                "forbidden": ("王晨",),
            },
        )

        # The negative-control conversation must not remain in the retained
        # front-end sample.  Cleanup is an acceptance condition for the user's
        # "only latest project / sources / conversation" requirement.
        negative_cleanup = _delete_records(client, project_id=None, conversation_ids=[negative_conversation_id], file_ids=[])
        negative_deleted = bool(negative_cleanup) and bool(negative_cleanup[0].get("deleted"))
        cleanup_step = run.step("negative-control conversation removed", component="Conversation API", expected="only the main conversation remains in a retained pass sample")
        run.observe(cleanup_step, negative_deleted, actual=json.dumps(negative_cleanup, ensure_ascii=False))
        if negative_deleted:
            negative_conversation_id = None

        evidence_path = run.write_evidence("project-memory-e2e.json", {
            "scenario": SCENARIO,
            "project_id": project_id,
            "main_conversation_id": primary_conversation_id,
            "source_file_ids": file_ids,
            "turn_count": len(scenario.turns),
            "turns": scenario.turns,
            "api_calls": client.calls,
        })
        for step in run.steps:
            step.evidence.setdefault("scenario_evidence", evidence_path)

        candidate_pass = not any(step.status == "FAIL" for step in run.steps)
        if candidate_pass:
            previous = _read_state()
            retention_step = run.step("replace only prior retained passed sample", component="Project/Conversation/Library APIs", expected="current sample retained and any verified prior CTX-PROJECT-01 sample removed")
            retention_actual: dict[str, Any] = {"previous_state_found": bool(previous), "current_retained": False}
            if previous:
                if not _owned_previous_state(client, previous):
                    run.observe(retention_step, False, actual=json.dumps({**retention_actual, "reason": "previous state failed ownership verification"}, ensure_ascii=False))
                else:
                    previous_outcomes = _delete_records(
                        client,
                        project_id=str(previous["project_id"]),
                        conversation_ids=[str(previous["conversation_id"])],
                        file_ids=[str(value) for value in previous["file_ids"]],
                    )
                    previous_deleted = all(bool(item.get("deleted")) for item in previous_outcomes)
                    run.observe(retention_step, previous_deleted, actual=json.dumps({**retention_actual, "previous_cleanup": previous_outcomes}, ensure_ascii=False))
            else:
                run.observe(retention_step, True, actual=json.dumps(retention_actual, ensure_ascii=False))
            if not any(step.status == "FAIL" for step in run.steps):
                _write_state(project_id=project_id, conversation_id=primary_conversation_id, file_ids=file_ids)
                retained = True
                run.notes.append("LATEST_PASSED_RETAINED: the project, its three sources, and only its main conversation remain for frontend inspection.")
        return "LCT_PASS" if retained else "TEST_FAIL"
    finally:
        if not retained:
            cleanup = _delete_records(client, project_id=project_id, conversation_ids=[value for value in (negative_conversation_id, primary_conversation_id) if value], file_ids=file_ids)
            if cleanup:
                run.write_evidence("cleanup.json", {"retained": False, "cleanup": cleanup})
            # A stale local state must not later authorize a deletion when its
            # associated retained sample was only partially removed elsewhere.
            current = _read_state()
            if current and current.get("project_id") == project_id:
                _remove_state()


def main() -> int:
    parser = argparse.ArgumentParser(description="Run CTX-PROJECT-01 real project/document long-memory scenario")
    parser.add_argument("--base-url", default="http://localhost:8000/api")
    parser.add_argument("--request-timeout-seconds", type=int, default=180)
    parser.add_argument("--index-wait-seconds", type=int, default=180)
    args = parser.parse_args()
    if args.request_timeout_seconds < 20:
        parser.error("--request-timeout-seconds must be at least 20")
    if args.index_wait_seconds < 3:
        parser.error("--index-wait-seconds must be at least 3")
    run = LctRun(lct=SCENARIO, repo_root=repository_root(), execution_path="LIVE_PROJECT_API + REAL_DOCX + REAL_MODEL + OWNER_SCOPED_CLEANUP")
    try:
        verdict = _run(args, run)
        return run.finish(lct_verdict=verdict, classification="real project/document long-conversation memory scenario")
    except Exception as exc:
        if isinstance(exc, ScenarioStop):
            run.notes.append("TERMINAL_TRANSPORT_BOUNDARY: " + str(exc))
        return cli_failure(run, exc, classification="real project/document long-conversation memory scenario")


if __name__ == "__main__":
    raise SystemExit(main())
