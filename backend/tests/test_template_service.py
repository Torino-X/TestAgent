from __future__ import annotations

import io
import hashlib
import zipfile
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError

from app.core.exceptions import (
    ForbiddenError,
    TemplateNotPublicError,
    TemplatePublishedDeleteConflictError,
    TemplateUnsupportedFileTypeError,
)
from app.models.conversation import Conversation
from app.models.template_asset import TemplateAsset
from app.models.template_version import TemplateVersion
from app.models.uploaded_file import UploadedFile
from app.models.user import User
from app.models.user_template import UserTemplate
from app.services.template_market_service import TemplateMarketService
from app.services.document_preview_service import DocumentPreviewService
from app.services.template_cover_service import TemplateCoverService
from app.services.template_preview_service import TemplatePreviewService
from app.services.template_service import TemplateService
from app.services.template_use_service import TemplateUseService
from app.utils.datetime import utcnow


class InMemoryTemplateStorage:
    storage_type = "memory"

    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}
        self._sequence = 0

    async def save_template(
        self,
        content: bytes,
        file_name: str,
        owner_id: str,
        template_id: str,
        version_no: int,
    ):
        self._sequence += 1
        key = f"templates/{owner_id}/{template_id}/v{version_no}/{self._sequence}_{file_name}"
        self.objects[key] = content
        return SimpleNamespace(storage_path=key, file_name=file_name, file_size=len(content))

    async def save_upload(self, content: bytes, original_name: str, user_id: str, conversation_id: str):
        self._sequence += 1
        key = f"uploads/{user_id}/{conversation_id}/{self._sequence}_{original_name}"
        self.objects[key] = content
        return SimpleNamespace(storage_path=key, file_name=original_name, file_size=len(content))

    async def save_preview(self, content: bytes, file_name: str, user_id: str, item_id: str):
        self._sequence += 1
        key = f"previews/{user_id}/{item_id}/{self._sequence}_{file_name}"
        self.objects[key] = content
        return SimpleNamespace(storage_path=key, file_name=file_name, file_size=len(content))

    async def read_bytes(self, storage_path: str) -> bytes:
        return self.objects[storage_path]

    async def open_file(self, storage_path: str):
        yield self.objects[storage_path]

    async def exists(self, storage_path: str) -> bool:
        return storage_path in self.objects

    async def delete_file(self, storage_path: str) -> None:
        self.objects.pop(storage_path, None)


def _docx_bytes(extra_text: str = "") -> bytes:
    document = Document()
    document.add_heading("标准测试方案", level=1)
    document.add_paragraph("测试范围与验收标准")
    if extra_text:
        document.add_paragraph(extra_text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(
            "xl/workbook.xml",
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="用例" sheetId="1" r:id="rId1"/></sheets></workbook>',
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>',
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>用例编号</t></is></c>'
            '<c r="B1" t="inlineStr"><is><t>预期结果</t></is></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>TC-001</t></is></c>'
            '<c r="B2" t="inlineStr"><is><t>通过</t></is></c></row></sheetData></worksheet>',
        )
    return output.getvalue()


@pytest.fixture(autouse=True)
def _fast_template_cover_renderer(monkeypatch):
    calls: list[int] = []

    def render(_service, content: bytes) -> bytes:
        calls.append(len(content))
        return b"%PDF-1.4\n% persisted template cover\n"

    monkeypatch.setattr(DocumentPreviewService, "_docx_to_pdf", render)
    return calls


async def _create_user(session, user_id: int, public_id: str) -> User:
    now = datetime(2026, 9, 15, 12, 0, 0)
    user = User(
        id=user_id,
        public_id=public_id,
        username=public_id,
        password_hash="test",
        role="user",
        status="active",
        created_at=now,
        updated_at=now,
    )
    session.add(user)
    await session.flush()
    return user


@pytest.mark.asyncio
async def test_template_upload_owner_library_created(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        content = _docx_bytes()

        result = await TemplateService(session, storage=storage).create_template(
            content=content,
            file_name="标准测试方案.docx",
            user_id=1,
            name="标准测试方案",
            category_code="test_plan",
            description="适用于通用 Web 项目",
            tags=["Web", "通用"],
            publish=False,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        await session.commit()

        asset = await session.get(TemplateAsset, 1)
        version = await session.get(TemplateVersion, 1)
        library_item = await session.get(UserTemplate, 1)
        assert result["source_type"] == "owner"
        assert asset is not None and asset.current_version_id == version.id
        assert version is not None and version.version_no == 1
        assert library_item is not None and library_item.version_id == version.id
        assert storage.objects[version.storage_path] == content


@pytest.mark.asyncio
async def test_template_upload_persists_cover_once_and_future_reads_never_convert(
    sqlite_session_factory,
    _fast_template_cover_renderer,
):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="持久封面.docx", user_id=1,
            name="持久封面", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        await session.commit()

        version = await session.get(TemplateVersion, 1)
        assert version is not None
        assert version.cover_status == "ready"
        assert version.cover_kind == "pdf"
        assert version.cover_storage_path in storage.objects
        assert len(_fast_template_cover_renderer) == 1

        market = await TemplateMarketService(session).list_market(1)
        assert market["items"][0]["cover"] == {
            "status": "ready",
            "kind": "pdf",
            "url": f"/api/templates/market/{item['template_id']}/preview/pdf",
            "spreadsheet": None,
        }

        preview_service = TemplatePreviewService(session, storage=storage)
        for _ in range(2):
            stream, _filename = await preview_service.get_pdf_stream(item["template_id"])
            assert b"".join([chunk async for chunk in stream]).startswith(b"%PDF")
        assert len(_fast_template_cover_renderer) == 1


@pytest.mark.asyncio
async def test_xlsx_upload_persists_bounded_first_sheet_cover(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        item = await TemplateService(session, storage=storage).create_template(
            content=_xlsx_bytes(), file_name="用例.xlsx", user_id=1,
            name="用例", category_code="test_case", description=None,
            tags=[], publish=True, content_type=None,
        )
        await session.commit()

        version = await session.get(TemplateVersion, 1)
        assert version is not None and version.cover_status == "ready"
        assert version.cover_kind == "spreadsheet"
        assert version.cover_payload_json["sheet_name"] == "用例"
        assert version.cover_payload_json["rows"][1] == ["TC-001", "通过"]
        preview = await TemplatePreviewService(session, storage=storage).get_preview(
            item["template_id"]
        )
        assert preview["spreadsheet"]["rows"][0] == ["用例编号", "预期结果"]


@pytest.mark.asyncio
async def test_legacy_missing_cover_backfills_once_and_stays_ready(
    sqlite_session_factory,
    _fast_template_cover_renderer,
):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="历史模板.docx", user_id=1,
            name="历史模板", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        await session.commit()
        version = await session.get(TemplateVersion, 1)
        old_cover_path = version.cover_storage_path
        await storage.delete_file(old_cover_path)
        version.cover_status = "missing"
        version.cover_kind = None
        version.cover_storage_path = None
        version.cover_size = None
        version.cover_generated_at = None
        await session.commit()

        assert len(_fast_template_cover_renderer) == 1
        cover_service = TemplateCoverService(session, storage=storage)
        await cover_service.backfill_missing_batch()
        await session.refresh(version)
        assert version.cover_status == "ready"
        assert version.cover_storage_path in storage.objects
        assert len(_fast_template_cover_renderer) == 2

        await cover_service.backfill_missing_batch()
        assert len(_fast_template_cover_renderer) == 2


@pytest.mark.asyncio
async def test_save_is_idempotent_and_save_count_only_increments_once(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_reader")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(),
            file_name="公开模板.docx",
            user_id=1,
            name="公开模板",
            category_code="test_plan",
            description=None,
            tags=[],
            publish=True,
            content_type=None,
        )

        market = TemplateMarketService(session)
        first = await market.save(owner_item["template_id"], 2)
        second = await market.save(owner_item["template_id"], 2)
        await session.commit()

        asset = await session.get(TemplateAsset, 1)
        assert first["already_saved"] is False
        assert second["already_saved"] is True
        assert first["user_template"]["id"] == second["user_template"]["id"]
        assert asset is not None and asset.save_count == 1


@pytest.mark.asyncio
async def test_market_saved_template_cannot_publish(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_reader")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(),
            file_name="公开模板.docx",
            user_id=1,
            name="公开模板",
            category_code="test_plan",
            description=None,
            tags=[],
            publish=True,
            content_type=None,
        )
        saved = await TemplateMarketService(session).save(owner_item["template_id"], 2)

        with pytest.raises(ForbiddenError):
            await TemplateService(session, storage=storage).publish(
                saved["user_template"]["id"], 2
            )


@pytest.mark.asyncio
async def test_use_template_new_conversation_materializes_independent_uploaded_file(
    sqlite_session_factory,
):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(),
            file_name="生成方案.docx",
            user_id=1,
            name="生成方案",
            category_code="test_plan",
            description=None,
            tags=[],
            publish=False,
            content_type=None,
        )

        used = await TemplateUseService(session, storage=storage).use(
            user_id=1,
            user_template_public_id=owner_item["id"],
            conversation_public_id=None,
        )
        await session.commit()

        conversation = await session.get(Conversation, 1)
        version = await session.get(TemplateVersion, 1)
        assert conversation is not None
        assert used["conversation"]["id"] == conversation.public_id
        assert used["uploaded_file"]["original_name"] == "生成方案.docx"
        uploaded_file = await session.get(UploadedFile, 1)
        assert uploaded_file is not None
        assert uploaded_file.storage_path != version.storage_path
        assert uploaded_file.storage_path in storage.objects
        assert "storage_path" not in used["uploaded_file"]


@pytest.mark.asyncio
async def test_use_template_other_users_library_forbidden(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_other")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(),
            file_name="私有模板.docx",
            user_id=1,
            name="私有模板",
            category_code="test_plan",
            description=None,
            tags=[],
            publish=False,
            content_type=None,
        )

        with pytest.raises(ForbiddenError):
            await TemplateUseService(session, storage=storage).use(
                user_id=2,
                user_template_public_id=owner_item["id"],
            )


@pytest.mark.asyncio
async def test_template_upload_rejects_unsupported_extension(sqlite_session_factory):
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        with pytest.raises(TemplateUnsupportedFileTypeError):
            await TemplateService(session, storage=InMemoryTemplateStorage()).create_template(
                content=b"plain text",
                file_name="unsafe.txt",
                user_id=1,
                name="Unsafe",
                category_code="test_plan",
                description=None,
                tags=[],
                publish=False,
                content_type="text/plain",
            )


@pytest.mark.asyncio
async def test_market_only_returns_public_active_and_honors_filters(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        service = TemplateService(session, storage=storage)
        await service.create_template(
            content=_docx_bytes(), file_name="公开方案.docx", user_id=1,
            name="Web 测试方案", category_code="test_plan", description="Web",
            tags=["Web"], publish=True, content_type=None,
        )
        await service.create_template(
            content=_docx_bytes(), file_name="私有用例.docx", user_id=1,
            name="私有测试用例", category_code="test_case", description="private",
            tags=[], publish=False, content_type=None,
        )

        result = await TemplateMarketService(session).list_market(
            1, query="Web", category="test_plan", sort="newest", page=1, page_size=10
        )

        assert result["total"] == 1
        assert result["items"][0]["name"] == "Web 测试方案"
        assert result["items"][0]["is_saved"] is True
        assert result["items"][0]["published_at"] is not None
        assert "storage_path" not in result["items"][0]


@pytest.mark.asyncio
async def test_mine_is_user_scoped(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_one")
        await _create_user(session, 2, "user_two")
        for user_id, label in ((1, "一号"), (2, "二号")):
            await TemplateService(session, storage=storage).create_template(
                content=_docx_bytes(), file_name=f"{label}.docx", user_id=user_id,
                name=f"{label}模板", category_code="test_plan", description=None,
                tags=[], publish=False, content_type=None,
            )

        one = await TemplateMarketService(session).list_mine(1)
        two = await TemplateMarketService(session).list_mine(2)

        assert [item["name"] for item in one["items"]] == ["一号模板"]
        assert [item["name"] for item in two["items"]] == ["二号模板"]


@pytest.mark.asyncio
async def test_market_preview_public_active_only_without_creating_uploaded_file(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        service = TemplateService(session, storage=storage)
        public_item = await service.create_template(
            content=_docx_bytes(), file_name="公开.docx", user_id=1,
            name="公开", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        private_item = await service.create_template(
            content=_docx_bytes(), file_name="私有.docx", user_id=1,
            name="私有", category_code="test_plan", description=None,
            tags=[], publish=False, content_type=None,
        )

        preview = await TemplatePreviewService(session, storage=storage).get_preview(
            public_item["template_id"]
        )
        assert preview["viewer"] == "word"
        assert preview["item"]["source"] == "template"
        assert await session.scalar(select(func.count(UploadedFile.id))) == 0
        with pytest.raises(TemplateNotPublicError):
            await TemplatePreviewService(session, storage=storage).get_preview(
                private_item["template_id"]
            )


@pytest.mark.asyncio
async def test_my_template_preview_uses_owned_pinned_version(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="私有方案.docx", user_id=1,
            name="私有方案", category_code="test_plan", description=None,
            tags=[], publish=False, content_type=None,
        )

        preview = await TemplatePreviewService(session, storage=storage).get_user_preview(
            item["id"], 1
        )

        assert preview["viewer"] == "word"
        assert preview["pdf_url"] == f"/api/templates/mine/{item['id']}/preview/pdf"
        listed = await TemplateMarketService(session).list_mine(1)
        assert listed["items"][0]["preview_url"] == f"/api/templates/mine/{item['id']}/preview"


@pytest.mark.asyncio
async def test_my_template_download_is_owner_scoped_and_creates_no_uploaded_file(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_other")
        item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="下载模板.docx", user_id=1,
            name="下载模板", category_code="test_plan", description=None,
            tags=[], publish=False, content_type=None,
        )

        stream, filename, media_type = await TemplateService(
            session, storage=storage
        ).get_download_stream(item["id"], 1)
        downloaded = b"".join([chunk async for chunk in stream])

        assert downloaded
        assert filename == "下载模板.docx"
        assert "wordprocessingml" in media_type
        assert await session.scalar(select(func.count(UploadedFile.id))) == 0
        with pytest.raises(ForbiddenError):
            await TemplateService(session, storage=storage).get_download_stream(item["id"], 2)


@pytest.mark.asyncio
async def test_remove_market_saved_and_published_owner_delete_conflict(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_reader")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="公开.docx", user_id=1,
            name="公开", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        saved = await TemplateMarketService(session).save(owner_item["template_id"], 2)

        await TemplateService(session, storage=storage).remove(saved["user_template"]["id"], 2)
        assert (await TemplateMarketService(session).list_mine(2))["total"] == 0
        with pytest.raises(TemplatePublishedDeleteConflictError):
            await TemplateService(session, storage=storage).remove(owner_item["id"], 1)


@pytest.mark.asyncio
async def test_use_template_existing_conversation_and_other_owner_forbidden(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_other")
        item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="会话模板.docx", user_id=1,
            name="会话模板", category_code="test_plan", description=None,
            tags=[], publish=False, content_type=None,
        )
        now = utcnow()
        session.add(Conversation(
            id=10, public_id="conv_owned", user_id=1, title="已有会话",
            status="active", created_at=now, updated_at=now,
        ))
        session.add(Conversation(
            id=20, public_id="conv_other", user_id=2, title="他人会话",
            status="active", created_at=now, updated_at=now,
        ))
        await session.flush()

        used = await TemplateUseService(session, storage=storage).use(
            user_id=1, user_template_public_id=item["id"], conversation_public_id="conv_owned"
        )
        assert used["conversation"]["id"] == "conv_owned"
        with pytest.raises(ForbiddenError):
            await TemplateUseService(session, storage=storage).use(
                user_id=1, user_template_public_id=item["id"], conversation_public_id="conv_other"
            )


@pytest.mark.asyncio
async def test_public_default_test_plan_template_materializes_as_confirmed_conversation_file(
    sqlite_session_factory,
):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "template_owner")
        await _create_user(session, 2, "requester")
        await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(),
            file_name="00_PlanWise_QA_测试方案模板.docx",
            user_id=1,
            name="00_PlanWise_QA_测试方案模板",
            category_code="test_plan",
            description=None,
            tags=[],
            publish=True,
            content_type=None,
        )
        now = utcnow()
        session.add(Conversation(
            id=30, public_id="conv_default_template", user_id=2, title="默认模板任务",
            status="active", created_at=now, updated_at=now,
        ))
        await session.flush()

        used = await TemplateUseService(session, storage=storage).use_public_default_test_plan_template(
            user_id=2,
            conversation_public_id="conv_default_template",
            template_name="00_PlanWise_QA_测试方案模板",
        )
        await session.flush()

        copied = (
            await session.execute(
                select(UploadedFile).where(
                    UploadedFile.public_id == used["uploaded_file"]["file_id"]
                )
            )
        ).scalar_one()
        assert copied is not None
        assert copied.conversation_id == 30
        assert copied.file_type == "test_plan_template"
        assert copied.storage_path in storage.objects
        assert (await session.execute(select(UserTemplate).where(UserTemplate.user_id == 2))).scalars().all() == []


@pytest.mark.asyncio
async def test_template_upload_accepts_xlsx_and_rejects_macro_payload(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        xlsx = await TemplateService(session, storage=storage).create_template(
            content=_xlsx_bytes(), file_name="测试用例.xlsx", user_id=1,
            name="测试用例", category_code="test_case", description=None,
            tags=[], publish=False,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        assert xlsx["file_ext"] == "xlsx"

        macro = io.BytesIO(_docx_bytes())
        with zipfile.ZipFile(macro, "a") as archive:
            archive.writestr("word/vbaProject.bin", b"macro")
        with pytest.raises(TemplateUnsupportedFileTypeError):
            await TemplateService(session, storage=storage).create_template(
                content=macro.getvalue(), file_name="宏模板.docx", user_id=1,
                name="宏模板", category_code="test_plan", description=None,
                tags=[], publish=False, content_type=None,
            )


@pytest.mark.asyncio
async def test_unpublish_is_owner_only(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_reader")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="公开.docx", user_id=1,
            name="公开", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        saved = await TemplateMarketService(session).save(owner_item["template_id"], 2)
        with pytest.raises(ForbiddenError):
            await TemplateService(session, storage=storage).unpublish(
                saved["user_template"]["id"], 2
            )


@pytest.mark.asyncio
async def test_restore_market_save_does_not_increment_save_count_again(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_reader")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="恢复.docx", user_id=1,
            name="恢复", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        first = await TemplateMarketService(session).save(owner_item["template_id"], 2)
        await TemplateService(session, storage=storage).remove(first["user_template"]["id"], 2)
        restored = await TemplateMarketService(session).save(owner_item["template_id"], 2)
        asset = await session.get(TemplateAsset, 1)
        assert restored["already_saved"] is False
        assert restored["user_template"]["id"] == first["user_template"]["id"]
        assert asset is not None and asset.save_count == 1


@pytest.mark.asyncio
async def test_my_template_download_uses_pinned_version(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_reader")
        first_content = _docx_bytes()
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=first_content, file_name="v1.docx", user_id=1,
            name="版本模板", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        saved = await TemplateMarketService(session).save(owner_item["template_id"], 2)
        asset = await session.get(TemplateAsset, 1)
        second_content = _docx_bytes("第二版")
        stored = await storage.save_template(second_content, "v2.docx", "1", asset.public_id, 2)
        session.add(TemplateVersion(
            id=2, public_id="tplv_second", template_id=asset.id, version_no=2,
            original_filename="v2.docx", file_ext="docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_size=len(second_content), file_hash=hashlib.sha256(second_content).hexdigest(),
            storage_path=stored.storage_path, created_at=utcnow(),
        ))
        asset.current_version_id = 2
        await session.flush()

        stream, filename, _mime = await TemplateService(
            session, storage=storage
        ).get_download_stream(saved["user_template"]["id"], 2)
        downloaded = b"".join([chunk async for chunk in stream])
        assert filename == "v1.docx"
        assert downloaded == first_content


@pytest.mark.asyncio
async def test_concurrent_save_unique_conflict_recovers_idempotently(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        await _create_user(session, 2, "user_reader")
        owner_item = await TemplateService(session, storage=storage).create_template(
            content=_docx_bytes(), file_name="并发.docx", user_id=1,
            name="并发", category_code="test_plan", description=None,
            tags=[], publish=True, content_type=None,
        )
        first = await TemplateMarketService(session).save(owner_item["template_id"], 2)
        existing = await session.scalar(select(UserTemplate).where(UserTemplate.user_id == 2))
        service = TemplateMarketService(session)
        service._user_templates.get_by_user_and_template = AsyncMock(
            side_effect=[None, existing]
        )
        service._user_templates.create = AsyncMock(
            side_effect=IntegrityError("insert user_templates", {}, RuntimeError("duplicate"))
        )

        raced = await service.save(owner_item["template_id"], 2)
        asset = await session.get(TemplateAsset, 1)
        assert raced["already_saved"] is True
        assert raced["user_template"]["id"] == first["user_template"]["id"]
        assert asset is not None and asset.save_count == 1


@pytest.mark.asyncio
async def test_template_upload_failure_cleans_canonical_object(sqlite_session_factory):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        service = TemplateService(session, storage=storage)
        service._versions.create = AsyncMock(side_effect=RuntimeError("database write failed"))

        with pytest.raises(RuntimeError, match="database write failed"):
            await service.create_template(
                content=_docx_bytes(), file_name="清理.docx", user_id=1,
                name="清理", category_code="test_plan", description=None,
                tags=[], publish=False, content_type=None,
            )
        assert storage.objects == {}


@pytest.mark.asyncio
async def test_template_upload_late_failure_cleans_cover_and_canonical_objects(
    sqlite_session_factory,
):
    storage = InMemoryTemplateStorage()
    async with sqlite_session_factory() as session:
        await _create_user(session, 1, "user_owner")
        service = TemplateService(session, storage=storage)
        service._user_templates.create = AsyncMock(
            side_effect=RuntimeError("owner library write failed")
        )

        with pytest.raises(RuntimeError, match="owner library write failed"):
            await service.create_template(
                content=_docx_bytes(), file_name="清理封面.docx", user_id=1,
                name="清理封面", category_code="test_plan", description=None,
                tags=[], publish=False, content_type=None,
            )
        assert storage.objects == {}
