from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from sqlalchemy.exc import OperationalError

from app.schemas.library import LibraryItem
from app.core.exceptions import DocumentPreviewEngineUnavailableError
from app.models.document_preview import DocumentPreview
from app.services.document_preview_service import DocumentPreviewService


def _stub_office_renderer(monkeypatch):
    captured: dict[str, object] = {}

    def fake_run(command, **_kwargs):
        source_path = Path(command[-1])
        output_path = Path(command[command.index("--outdir") + 1])
        captured["command"] = command
        captured["source"] = source_path.read_bytes()
        (output_path / "document.pdf").write_bytes(b"%PDF-1.7\nfull-layout-preview")
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(DocumentPreviewService, "_office_binary", staticmethod(lambda: "soffice"))
    monkeypatch.setattr("app.services.document_preview_service.subprocess.run", fake_run)
    return captured


def _xlsx_with_one_sheet() -> bytes:
    parts = {
        "xl/workbook.xml": """<workbook xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\" xmlns:r=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships\"><sheets><sheet name=\"测试表\" sheetId=\"1\" r:id=\"rId1\"/></sheets></workbook>""",
        "xl/_rels/workbook.xml.rels": """<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"><Relationship Id=\"rId1\" Target=\"worksheets/sheet1.xml\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet\"/></Relationships>""",
        "xl/sharedStrings.xml": """<sst xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\"><si><t>项目名称</t></si><si><t>测试环境</t></si></sst>""",
        "xl/worksheets/sheet1.xml": """<worksheet xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\"><sheetData><row r=\"1\"><c r=\"A1\" t=\"s\"><v>0</v></c><c r=\"B1\"><v>42</v></c></row><row r=\"2\"><c r=\"A2\" t=\"s\"><v>1</v></c></row></sheetData></worksheet>""",
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for path, text in parts.items():
            archive.writestr(path, text)
    return output.getvalue()


def test_xlsx_preview_reads_real_cells_without_office_runtime():
    preview = DocumentPreviewService._parse_spreadsheet(DocumentPreviewService.__new__(DocumentPreviewService), _xlsx_with_one_sheet(), ".xlsx")

    assert preview.sheet_name == "测试表"
    assert preview.columns == ["A", "B"]
    assert preview.rows == [["项目名称", "42"], ["测试环境", ""]]


def test_docx_preview_delegates_original_ooxml_to_office_renderer(monkeypatch):
    document = Document()
    document.add_heading("测试方案", level=1)
    document.add_paragraph("这是用于预览的文档内容。")
    raw = io.BytesIO()
    document.save(raw)

    preview = DocumentPreviewService.__new__(DocumentPreviewService)
    captured = _stub_office_renderer(monkeypatch)
    pdf = preview._docx_to_pdf(raw.getvalue())

    assert pdf == b"%PDF-1.7\nfull-layout-preview"
    assert b"word/document.xml" in captured["source"]
    assert "pdf:writer_pdf_Export" in captured["command"]
    assert any(str(argument).startswith("-env:UserInstallation=file:") for argument in captured["command"])


def test_docx_preview_uses_native_word_exporter_when_word_is_available(monkeypatch):
    preview = DocumentPreviewService.__new__(DocumentPreviewService)
    monkeypatch.setattr(preview, "_office_binary", lambda: r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE")
    monkeypatch.setattr(preview, "_docx_to_pdf_with_word", lambda content: b"%PDF-word-native\n" + content[:4])

    assert preview._docx_to_pdf(b"PK\x03\x04document") == b"%PDF-word-native\nPK\x03\x04"


def test_word_exporter_without_desktop_access_has_a_specific_preview_error(monkeypatch):
    preview = DocumentPreviewService.__new__(DocumentPreviewService)
    monkeypatch.setattr(preview, "_office_binary", lambda: r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE")
    monkeypatch.setattr(preview, "_docx_to_pdf_with_word", lambda _content: (_ for _ in ()).throw(RuntimeError("desktop session unavailable")))

    with pytest.raises(DocumentPreviewEngineUnavailableError):
        preview._docx_to_pdf(b"PK\x03\x04document")


def test_missing_office_engine_has_a_specific_preview_error(monkeypatch):
    monkeypatch.setattr("app.services.document_preview_service.shutil.which", lambda _name: None)
    monkeypatch.setattr("app.services.document_preview_service.get_settings", lambda: SimpleNamespace(document_preview_office_bin=""))
    monkeypatch.setattr("app.services.document_preview_service.Path.is_file", lambda _path: False)

    with pytest.raises(DocumentPreviewEngineUnavailableError):
        DocumentPreviewService._office_binary()


@pytest.mark.asyncio
async def test_dotless_docx_artifact_extension_uses_word_preview():
    item = LibraryItem(
        id="art_docx",
        name="generated-plan.docx",
        kind="file",
        source="generated",
        extension="docx",
        modified_at=datetime.now(timezone.utc),
    )

    class FakeLibrary:
        async def get_owned_storage_item(self, item_id: str, user_id: int):
            return SimpleNamespace(storage_path="artifacts/generated-plan.docx")

        @staticmethod
        def _to_item(_item):
            return item

    preview = DocumentPreviewService.__new__(DocumentPreviewService)
    preview._library = FakeLibrary()
    preview._request_word_preview = AsyncMock(return_value={"status": "queued", "source_version": "hash:abc", "error": None})

    result = await preview.get_preview("art_docx", 1)

    assert result["viewer"] == "word"
    assert result["pdf_url"] is None
    assert result["preview_status"] == "queued"


@pytest.mark.asyncio
async def test_dotless_docx_artifact_extension_converts_to_preview_pdf(monkeypatch):
    document = Document()
    document.add_paragraph("generated preview")
    raw = io.BytesIO()
    document.save(raw)

    class FakeLibrary:
        async def get_owned_storage_item(self, item_id: str, user_id: int):
            return SimpleNamespace(storage_path="artifacts/generated-plan.docx", file_ext="docx")

        @staticmethod
        def _item_name(_item):
            return "generated-plan.docx"

    preview = DocumentPreviewService.__new__(DocumentPreviewService)
    preview._library = FakeLibrary()
    preview._get_ready_preview = AsyncMock(return_value=SimpleNamespace(preview_storage_path="previews/art_docx.pdf"))
    preview._read_bytes = AsyncMock(return_value=b"%PDF-1.7\ncached-preview")

    pdf, filename = await preview.get_pdf_bytes("art_docx", 1)

    assert filename == "generated-plan.pdf"
    assert pdf == b"%PDF-1.7\ncached-preview"
    preview._read_bytes.assert_awaited_once_with("previews/art_docx.pdf")


def test_source_version_uses_content_hash_before_storage_path():
    service = DocumentPreviewService.__new__(DocumentPreviewService)
    original = SimpleNamespace(file_hash="sha256:abc", version_no=1, storage_path="uploads/one.docx")
    renamed = SimpleNamespace(file_hash="sha256:abc", version_no=1, storage_path="uploads/one.docx")

    assert service._source_version(original) == service._source_version(renamed) == "hash:sha256:abc"


def test_preview_model_eagerly_loads_server_generated_timestamps():
    """Status serialization must not issue synchronous ORM IO after flush."""
    assert DocumentPreview.__mapper__.eager_defaults is True


def test_processing_preview_is_recovered_after_office_timeout(monkeypatch):
    monkeypatch.setattr(
        "app.services.document_preview_service.get_settings",
        lambda: SimpleNamespace(document_preview_office_timeout_seconds=90),
    )
    stale = SimpleNamespace(updated_at=datetime.now(timezone.utc).replace(year=2020))
    current = SimpleNamespace(updated_at=datetime.now(timezone.utc))

    assert DocumentPreviewService._processing_is_stale(stale) is True
    assert DocumentPreviewService._processing_is_stale(current) is False


def test_cached_processing_preview_recovers_after_office_timeout(monkeypatch):
    monkeypatch.setattr(
        "app.services.document_preview_service.get_settings",
        lambda: SimpleNamespace(document_preview_office_timeout_seconds=90),
    )

    assert DocumentPreviewService._cached_processing_is_stale({"updated_at": "2020-01-01T00:00:00+00:00"}) is True
    assert DocumentPreviewService._cached_processing_is_stale({"updated_at": "not-a-date"}) is False


@pytest.mark.asyncio
async def test_queued_cache_without_durable_preview_recreates_and_schedules_job(monkeypatch):
    """A Redis queued hint must not block recovery when its DB row was lost."""

    cache = SimpleNamespace(
        get=AsyncMock(return_value={
            "status": "queued",
            "source_version": "hash:preview-source",
            "error": None,
        }),
        set=AsyncMock(),
    )
    monkeypatch.setattr(
        "app.cache.domains.document_preview_cache.get_document_preview_cache",
        lambda: cache,
    )

    session = SimpleNamespace(add=lambda _preview: None, flush=AsyncMock())
    service = DocumentPreviewService.__new__(DocumentPreviewService)
    service._session = session
    service._preview_job_required = False
    service._find_preview = AsyncMock(return_value=None)

    result = await service._request_word_preview(
        "art_preview", 1, SimpleNamespace(file_hash="preview-source", version_no=1, storage_path="artifacts/test.docx")
    )

    assert result["status"] == "queued"
    assert service.preview_job_required is True
    service._find_preview.assert_awaited_once_with("art_preview", 1)
    session.flush.assert_awaited_once()
    cache.set.assert_awaited_once()


@pytest.mark.asyncio
async def test_preview_claim_lock_timeout_is_recoverable():
    """A competing DB transaction must not make the HTTP background task fail."""

    class MySqlLockTimeout(Exception):
        pass

    session = SimpleNamespace(
        execute=AsyncMock(
            side_effect=OperationalError(
                "UPDATE document_previews",
                {},
                MySqlLockTimeout(1205, "Lock wait timeout exceeded; try restarting transaction"),
            )
        ),
        rollback=AsyncMock(),
    )
    service = DocumentPreviewService.__new__(DocumentPreviewService)
    service._session = session

    await service._build_queued_preview("art_preview", 1)

    session.rollback.assert_awaited_once()
