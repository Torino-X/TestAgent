from app.tools.section_suggestion_tool import SectionSuggestionTool
from app.tools.template_parser_tool import TemplateParserTool
from app.common.template_section_service import TemplateSectionService
from docx import Document


def test_template_parser_attaches_canonical_identity_to_nested_section_tree():
    sections = [
        {
            "title": "1 项目概述",
            "level": 1,
            "mode": "ai",
            "children": [
                {
                    "title": "1.1 测试范围",
                    "level": 2,
                    "mode": "ai",
                    "children": [],
                }
            ],
        }
    ]
    bindings = [
        {
            "section_id": "body_10_level_1",
            "field": "overview",
            "suggested_field": "overview",
            "clean_title": "项目概述",
        },
        {
            "section_id": "body_12_level_2",
            "field": "test_scope",
            "suggested_field": "test_scope",
            "clean_title": "测试范围",
        },
    ]

    TemplateParserTool._attach_section_bindings(sections, bindings)

    suggestions = SectionSuggestionTool._build_suggestions(sections)
    assert [item["section_id"] for item in suggestions] == [
        "body_10_level_1",
        "body_12_level_2",
    ]
    assert sections[0]["field"] == "overview"
    assert sections[0]["children"][0]["field"] == "test_scope"


def test_numbered_requirement_sentences_are_not_misclassified_as_headings(tmp_path):
    template = Document()
    template.add_paragraph("1. 测试策略")
    template.add_paragraph("1. 明确需要重点关注或具有特殊性的测试内容。")
    template.add_paragraph("2. 测试工具二：服务于测试特性相关测试；如需开发，请补充以下需求：")
    template.add_paragraph("6. 测试用例规划")
    file_path = tmp_path / "numbered-template.docx"
    template.save(file_path)

    result = TemplateSectionService().parse_template(str(file_path))
    titles = [section.title for section in result.sections]

    assert titles == ["1. 测试策略", "6. 测试用例规划"]
